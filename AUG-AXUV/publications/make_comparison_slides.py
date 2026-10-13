# %%
"""Build 16:9 comparison slides (synthetic vs. experimental, linear, untitled plots).

Each slide holds a 2x2 grid: top row synthetic, bottom row experimental,
left column vertical camera, right column horizontal camera.
One PPTX/PDF pair is written per synthetic variant into $AXUV_OUTPUT_DIR.

Run with:
    uv run --with python-pptx python publications/make_comparison_slides.py
"""

import shutil
import subprocess
import tempfile
from datetime import date

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from axuv.io import OUTPUT_DIR

output_dir = OUTPUT_DIR

SHOTS = [40673, 41007]
# sector -> (injection position folder, vertical camera, horizontal camera)
SECTORS = {16: ("P1", "D16", "DHT"), 5: ("P45", "DVC", "DHC")}
SECTOR_ORDER = [16, 5]
REFLECTION = "refl_lowres_masked"
VARIANTS = ["nominal_no_noise", "degraded_no_noise"]

SLIDE_W, SLIDE_H = 13.333, 7.5  # inches, 16:9
TITLE_H = 0.55
COL_LABEL_H = 0.35
ROW_LABEL_W = 0.45
FOOTER_H = 0.3
AUTHOR = "Ferenc Lengyel"


def image_paths(shot, sector, variant):
    pos, vcam, hcam = SECTORS[sector]
    shot_dir = output_dir / str(shot)
    syn_dir = shot_dir / pos / REFLECTION / variant
    return {
        ("syn", "vert"): syn_dir / f"{shot}_S{sector}_vert_notitle_lin.png",
        ("syn", "horiz"): syn_dir / f"{shot}_S{sector}_horiz_notitle_lin.png",
        ("exp", "vert"): shot_dir / f"{shot}_{vcam}_S{sector}_vert_notitle_lin.png",
        ("exp", "horiz"): shot_dir / f"{shot}_{hcam}_S{sector}_horiz_notitle_lin.png",
    }


def add_text(slide, text, left, top, width, height, size, bold=False, rotation=0,
             align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.rotation = rotation
    frame = box.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = False
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return box


def add_slide(prs, shot, sector, variant, number):
    pos, vcam, hcam = SECTORS[sector]
    paths = image_paths(shot, sector, variant)
    for p in paths.values():
        if not p.exists():
            raise FileNotFoundError(p)

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Largest cell with the image aspect ratio that fits the free area
    w_px, h_px = Image.open(paths[("syn", "vert")]).size
    aspect = w_px / h_px
    avail_w = SLIDE_W - ROW_LABEL_W - 0.1
    avail_h = SLIDE_H - TITLE_H - COL_LABEL_H - FOOTER_H
    cell_h = min(avail_h / 2, avail_w / 2 / aspect)
    cell_w = cell_h * aspect

    grid_w = 2 * cell_w
    x0 = ROW_LABEL_W + (SLIDE_W - ROW_LABEL_W - grid_w) / 2
    y0 = TITLE_H + COL_LABEL_H

    add_text(
        slide,
        f"#{shot}, sector {sector} ({pos}), synthetic: {REFLECTION}/{variant}",
        0, 0.05, SLIDE_W, TITLE_H - 0.05, 22, bold=True,
    )

    cols = {"vert": f"Vertical ({vcam})", "horiz": f"Horizontal ({hcam})"}
    rows = {"syn": "Synthetic", "exp": "Experimental"}
    for j, (cam, label) in enumerate(cols.items()):
        add_text(slide, label, x0 + j * cell_w, TITLE_H, cell_w, COL_LABEL_H, 18, bold=True)
    for i, (kind, label) in enumerate(rows.items()):
        # Rotated box: width/height are swapped visually, so center it on the row
        cy = y0 + (i + 0.5) * cell_h
        cx = x0 - ROW_LABEL_W / 2
        add_text(slide, label, cx - cell_h / 2, cy - ROW_LABEL_W / 2, cell_h, ROW_LABEL_W,
                 18, bold=True, rotation=270)
        for j, cam in enumerate(cols):
            slide.shapes.add_picture(
                str(paths[(kind, cam)]),
                Inches(x0 + j * cell_w), Inches(y0 + i * cell_h),
                Inches(cell_w), Inches(cell_h),
            )

    footer_y = SLIDE_H - FOOTER_H
    footer_w = SLIDE_W / 3 - 0.3
    add_text(slide, AUTHOR, 0.3, footer_y, footer_w, FOOTER_H, 12, align=PP_ALIGN.LEFT)
    add_text(slide, date.today().isoformat(), SLIDE_W / 3, footer_y, SLIDE_W / 3, FOOTER_H, 12)
    add_text(slide, str(number), 2 * SLIDE_W / 3, footer_y, footer_w, FOOTER_H, 12,
             align=PP_ALIGN.RIGHT)


def build(variant):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slides = [(shot, sector) for shot in SHOTS for sector in SECTOR_ORDER]
    for number, (shot, sector) in enumerate(slides, start=1):
        add_slide(prs, shot, sector, variant, number)

    stem = f"comparison_slides_{REFLECTION}_{variant}"
    pptx_path = output_dir / f"{stem}.pptx"
    prs.save(pptx_path)

    # Separate LibreOffice profile so a running instance does not block conversion
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless",
             "--convert-to", "pdf", "--outdir", str(output_dir), str(pptx_path)],
            check=True, capture_output=True,
        )
    print(f"Wrote {pptx_path} and {pptx_path.with_suffix('.pdf')}")


# %%
if __name__ == "__main__":
    if shutil.which("soffice") is None:
        raise RuntimeError("soffice (LibreOffice) is required for PDF conversion")
    for variant in VARIANTS:
        build(variant)
